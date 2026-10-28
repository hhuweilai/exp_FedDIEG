from __future__ import annotations

import hashlib
import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


REFERENCE_DOCX = Path(os.environ["REFERENCE_DOCX"])
OUTPUT_DOCX = Path(os.environ["OUTPUT_DOCX"])


EXPECTED_REFERENCE_SHA256 = "75b5086b586f86d85f302bc836f9cae52b62ec0fe3a923570ca3c9bd23ff0be3"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def set_run_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "宋体"
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), "宋体")
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")


def set_paragraph_font(paragraph, size: float = 10.5) -> None:
    for run in paragraph.runs:
        set_run_font(run, size=size)


def set_cell_text(cell, text: str, bold: bool = False, size: float = 9.5) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text = ""
    parts = text.split("\n")
    for idx, part in enumerate(parts):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(part)
        set_run_font(run, size=size, bold=bold)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "666666")


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_col_widths(table, widths: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = Inches(width)
                tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
                tc_w = tc_pr.find(qn("w:tcW"))
                if tc_w is None:
                    tc_w = OxmlElement("w:tcW")
                    tc_pr.append(tc_w)
                tc_w.set(qn("w:w"), str(int(width * 1440)))
                tc_w.set(qn("w:type"), "dxa")


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.allow_autofit = False
    try:
        table.style = "Normal Table"
    except Exception:
        pass
    set_table_borders(table)
    mark_header_row(table.rows[0])
    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_text(header_cells[idx], header, bold=True, size=9.5)
        set_cell_shading(header_cells[idx], "EDEDED")
        set_cell_margins(header_cells[idx])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value, size=9.2)
            set_cell_margins(cells[idx])
    set_col_widths(table, widths)
    doc.add_paragraph()


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading(level=level)
    p.paragraph_format.space_before = Pt(6 if level == 1 else 3)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=14 if level == 1 else 12, bold=True)


def add_body(doc: Document, text: str, indent: bool = False) -> None:
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Pt(22)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.18
    run = p.add_run(text)
    set_run_font(run, size=10.5)


def add_subtopic(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title)
    set_run_font(run, size=10.5, bold=True)


def add_bullet(doc: Document, text: str) -> None:
    try:
        p = doc.add_paragraph(style="List Bullet")
    except Exception:
        p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run_font(run, size=10.5)


def clear_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def apply_base_style(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
    for style_name in ["Normal", "Normal (Web)", "List Paragraph", "List Bullet"]:
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = "宋体"
            style.font.size = Pt(10.5)
            r_pr = style._element.get_or_add_rPr()
            r_fonts = r_pr.rFonts
            if r_fonts is None:
                r_fonts = OxmlElement("w:rFonts")
                r_pr.append(r_fonts)
            r_fonts.set(qn("w:eastAsia"), "宋体")
            r_fonts.set(qn("w:ascii"), "Times New Roman")
            r_fonts.set(qn("w:hAnsi"), "Times New Roman")


def build_document() -> None:
    if not REFERENCE_DOCX.exists():
        raise FileNotFoundError(REFERENCE_DOCX)
    if sha256(REFERENCE_DOCX) != EXPECTED_REFERENCE_SHA256:
        raise RuntimeError("Reference DOCX changed; rerun template distillation before generating.")

    doc = Document(str(REFERENCE_DOCX))
    clear_body(doc)
    apply_base_style(doc)

    add_heading(doc, "项目名称：FedDIEG：面向动态兴趣演化的联邦序列推荐模型", 1)

    add_heading(doc, "项目核心词汇：", 2)
    add_table(
        doc,
        ["分类", "核心词汇"],
        [
            ["一、项目架构类", "FedDIEG、联邦序列推荐、SASRec、Transformer Encoder、客户端动态兴趣图、兴趣记忆库"],
            ["二、联邦推荐类", "Federated Recommendation、客户端采样、本地训练、参数/表示聚合、数据隔离、隐私风险"],
            ["三、序列建模类", "item embedding、position embedding、self-attention、causal mask、短期兴趣 summary、最后有效位置打分"],
            ["四、动态图与邻居检索类", "Top-K/KNN、余弦相似度、temperature softmax、min_sim 阈值、动态图边、边权重"],
            ["五、图融合与稳定性类", "图邻居消息聚合、sigmoid gate、LayerNorm、残差注入、时间一致性、结构一致性、edge EMA"],
            ["六、冷启动类", "新客户端、少量初始交互、原型语义锚定、老客户端记忆库、prototype score fusion"],
            ["七、训练评估类", "PyTorch、GPU microbatch、负采样、leave-one-out、候选集评估、HR@10、Recall@10、NDCG@10"],
            ["八、实验结果类", "MovieLens-1M、Amazon Digital Music、Video Games、Arts/Crafts/Sewing、消融实验、动态窗口、敏感性分析"],
        ],
        [1.65, 4.95],
    )

    add_heading(doc, "项目基本信息", 2)
    add_table(
        doc,
        ["项目", "信息"],
        [
            ["项目类型", "论文研究 / 算法设计 / 联邦推荐系统实验"],
            ["技术栈", "Python、PyTorch、SASRec/Transformer、动态图建模、推荐系统排序评估"],
            ["核心问题", "数据分散不能集中上传、用户兴趣随时间漂移、新客户端交互稀疏导致冷启动"],
            ["核心思路", "本地序列编码 + 服务器端兴趣记忆库 + Top-K 动态兴趣图 + 图门控融合 + 一致性约束 + 原型锚定"],
            ["数据集", "MovieLens-1M、Amazon Digital Music、Amazon Video Games、Amazon Arts/Crafts/Sewing"],
            ["主要指标", "HR@10、Recall@10、NDCG@10；动态实验还关注 avg_window_ndcg、boundary_drop、avg_ecr、stability_score"],
            ["对比方法", "FedMF、FedSASRec、GPFedRec、FCUCR、TarDGR、F3CRec、DGFedRS"],
            ["工程实现", "单机 GPU 批量化模拟联邦训练，每个用户模拟为一个客户端，统一负采样、候选集和评估协议"],
        ],
        [1.55, 5.05],
    )

    add_heading(doc, "项目介绍", 2)
    add_body(
        doc,
        "本项目是一个面向联邦推荐场景的动态序列推荐模型研究，主要解决传统推荐系统中用户行为数据集中训练带来的隐私风险、数据分散场景下的协同困难，以及用户兴趣随时间变化导致的推荐效果不稳定问题。",
        indent=True,
    )
    add_body(
        doc,
        "系统整体采用“本地序列建模与服务器端表示协同分离”的思路。每个用户被模拟为一个客户端，本地保留自己的交互序列；模型以 SASRec/Transformer 作为序列编码主干，提取用户当前兴趣表示；服务器侧维护客户端兴趣记忆库，并基于记忆库构建 Top-K 动态兴趣图，让兴趣相似的客户端在表示层面互相增强。",
        indent=True,
    )
    add_body(
        doc,
        "FedDIEG 的重点不是简单把 SASRec 放进联邦训练框架，而是在联邦设定下进一步建模客户端之间随时间变化的兴趣关系。项目中加入短期兴趣注意力、图邻居消息聚合、图门控融合、时间一致性、结构一致性和原型语义锚定机制，并在 MovieLens-1M 和 Amazon 多领域数据集上完成主实验、动态窗口实验、消融实验、冷启动实验和超参数敏感性分析。",
        indent=True,
    )

    add_heading(doc, "项目工作流程", 2)
    workflow = [
        ("数据准备与客户端划分", "项目先读取 MovieLens-1M 和 Amazon 多领域数据集，按 leave-one-out 协议划分训练集、验证集和测试集。实验中每个用户被视为一个联邦客户端，本地交互序列不集中上传，只在单机上模拟联邦训练过程。"),
        ("本地序列编码", "每轮训练采样一批参与客户端，把其本地交互序列转换为 item embedding，并加入位置编码，输入 SASRec/Transformer 进行序列建模，得到每个时间步的隐藏表示。"),
        ("短期兴趣提取", "FedDIEG 对 SASRec 输出的序列隐藏状态做 attention pooling，得到客户端当前阶段的短期兴趣 summary。该表示比单纯取最后一个物品更稳，也比全序列平均更关注近期兴趣。"),
        ("兴趣记忆库更新", "服务器端维护每个客户端的兴趣记忆向量。客户端参与训练后，用当前短期兴趣 summary 按动量方式更新记忆库，使兴趣状态在跨轮次训练中保持连续。"),
        ("Top-K 动态兴趣图构建", "对客户端记忆向量做 L2 归一化后计算余弦相似度，为每个客户端检索 Top-K 相似邻居，并用 temperature softmax 得到边权。低于 min_sim 阈值的邻居会被过滤。"),
        ("图邻居聚合与门控融合", "模型对邻居记忆做加权聚合，得到图增强向量，再根据本地短期兴趣和图增强向量计算 sigmoid gate，动态决定引入多少邻居信息，避免错误邻居造成负迁移。"),
        ("损失优化与稳定约束", "训练目标由推荐主损失、最终位置预测增强损失、时间一致性损失和结构一致性损失组成。时间一致性约束客户端兴趣表示不要剧烈跳变，结构一致性约束动态图边关系不要频繁震荡。"),
        ("统一评估与结果分析", "所有方法使用相同的数据划分、客户端采样、负采样、候选集和指标计算方式，输出 round_metrics、summary、window_metrics 等结果文件，支撑主实验、消融、动态窗口、冷启动和敏感性分析。"),
    ]
    for title, text in workflow:
        add_subtopic(doc, title)
        add_body(doc, text)

    add_heading(doc, "核心技术原理", 2)
    principles = [
        (
            "联邦推荐与序列推荐",
            "联邦推荐的核心是把用户原始行为数据保留在本地，服务器只处理模型参数、梯度或表示层面的信息。序列推荐的目标是根据用户按时间排序的历史交互预测下一个可能感兴趣的物品。FedDIEG 将二者结合起来：每个客户端本地进行序列建模，服务器侧只维护兴趣记忆和相似关系。面试时要强调，项目是公开数据集上的联邦模拟实验，不是真实多设备部署。",
        ),
        (
            "SASRec/Transformer 序列编码",
            "SASRec 用 self-attention 建模历史交互序列。输入序列先转成 item embedding，再加入位置编码，并使用 causal mask 防止当前位置看到未来交互。输出的序列隐藏状态可以用于训练正负样本打分，也可以取最后有效位置作为推荐表示。它比矩阵分解更适合捕捉用户行为顺序和短期兴趣变化。",
        ),
        (
            "短期兴趣注意力",
            "只取最后一个位置可能受单次异常行为影响，直接平均整个历史又会稀释近期兴趣。因此 FedDIEG 用可学习查询向量对序列隐藏状态做 attention pooling，得到 short-term interest summary。这个 summary 后续会参与记忆库更新、Top-K 邻居检索、图门控融合和最终打分。",
        ),
        (
            "客户端兴趣记忆库",
            "兴趣记忆库保存每个客户端的历史兴趣向量，可以理解为服务器端看到的表示级用户状态。更新时使用 EMA/动量形式：新的记忆等于历史记忆和当前短期兴趣的加权组合。这样既能反映新兴趣，又能降低单轮采样噪声导致的表示震荡。",
        ),
        (
            "Top-K 动态兴趣图",
            "动态兴趣图的节点是客户端，边表示客户端之间的兴趣相似关系。项目中对记忆向量做归一化，用余弦相似度选择 Top-K 邻居，并通过 temperature softmax 计算邻居权重。图是动态更新的，因为用户兴趣和客户端记忆会随训练轮次或时间窗口变化。",
        ),
        (
            "图邻居消息聚合与门控融合",
            "邻居信息不能无条件加入本地表示，否则可能出现负迁移。FedDIEG 先对邻居记忆做加权聚合，得到图增强向量，再用本地短期兴趣和图增强向量计算 gate。最终通过残差方式把 gate 控制后的图增强注入序列表示，且更偏向影响序列后部位置。",
        ),
        (
            "时间一致性与结构一致性",
            "时间一致性关注同一客户端自己的兴趣表示在相邻阶段不要无意义跳变；结构一致性关注客户端图的邻居边不要因为轻微噪声频繁变化。二者都是辅助约束，用 lambda_temp 和 lambda_struct 控制强度。约束太强会降低适应速度，约束太弱又难以稳定动态图。",
        ),
        (
            "原型语义锚定",
            "原型语义锚定用于新客户端冷启动。新客户端只有少量初始交互时，先提取一个初始兴趣表示，再到老客户端记忆库中检索相似原型，将老客户端稳定兴趣作为新客户端的语义锚点。它可以作用于记忆初始化，也可以在候选物品打分阶段融入原型分数。",
        ),
        (
            "GPU 批量化联邦训练",
            "传统联邦模拟逐客户端训练会非常慢。项目将同一轮采样到的客户端组织成 microbatch，在 GPU 上并行完成序列编码、负采样、邻居检索、图融合和损失计算。这样保留了按客户端采样的联邦逻辑，同时显著提高实验效率。",
        ),
    ]
    for title, text in principles:
        add_subtopic(doc, title)
        add_body(doc, text)

    add_heading(doc, "数据集与实验结果", 2)
    add_table(
        doc,
        ["数据集", "用户数", "物品数", "交互数", "稀疏度特点"],
        [
            ["MovieLens-1M", "6,040", "3,416", "999,611", "相对稠密，sparsity 约 95.16%"],
            ["Amazon Digital Music", "12,381", "9,906", "123,518", "高稀疏，sparsity 约 99.90%"],
            ["Amazon Video Games", "50,626", "16,882", "453,881", "高稀疏，sparsity 约 99.95%"],
            ["Amazon Arts/Crafts/Sewing", "45,346", "21,050", "392,935", "高稀疏，sparsity 约 99.96%"],
        ],
        [1.75, 0.8, 0.8, 0.95, 2.3],
    )
    add_table(
        doc,
        ["数据集", "FedDIEG NDCG@10", "最强基线", "基线 NDCG@10", "相对提升"],
        [
            ["Amazon Arts/Crafts/Sewing", "0.3662", "TarDGR", "0.3614", "+1.31%"],
            ["Amazon Digital Music", "0.3357", "TarDGR", "0.3247", "+3.37%"],
            ["MovieLens-1M", "0.4934", "TarDGR", "0.4828", "+2.18%"],
            ["Amazon Video Games", "0.4491", "GPFedRec", "0.4451", "+0.90%"],
        ],
        [1.9, 1.25, 1.1, 1.25, 1.1],
    )
    add_body(
        doc,
        "主实验采用统一的 leave-one-out 与候选集评估协议。面试时可以稳妥表述为：FedDIEG 在静态主实验中相对最强基线整体取得提升，其中 Digital Music 提升约 3.37%，MovieLens-1M 提升约 2.18%，不同数据集提升幅度存在差异。",
    )
    add_table(
        doc,
        ["实验", "关键指标", "结论口径"],
        [
            ["动态窗口实验", "avg_window_ndcg、final_window_ndcg、adaptation_gain、boundary_drop", "Digital Music 和 MovieLens-1M 上 FedDIEG 的 avg_window_ndcg 优于最强基线，Video Games 和 Arts/Crafts 与最强基线接近但略低。"],
            ["消融实验", "avg_ecr、stability_score、max_drop、avg_window_ndcg", "在 Arts/Crafts 多随机种子结果中，Full 将 avg_ecr 从 Plain 的约 0.9287 降到约 0.0129，稳定性得分从约 0.2228 提升到约 0.2662。"],
            ["冷启动实验", "initial_ndcg、auc_ndcg、final_ndcg", "FedDIEG 在四个数据集平均 initial_ndcg 和 auc_ndcg 上表现最好；NoPrototype 会降低平均 AUC，说明原型语义锚定有贡献。"],
            ["敏感性分析", "K、rho、lambda_temp、lambda_struct", "K=20/40 往往带来更高窗口 NDCG；rho 对边变化率影响明显；lambda_temp 和 lambda_struct 体现精度与稳定性的权衡。"],
        ],
        [1.25, 1.65, 3.7],
    )

    add_heading(doc, "关键模块逐文件说明", 1)
    add_body(
        doc,
        "这一节用于把仓库文件和面试回答对应起来。面试官如果问“这个文件是干什么的”，可以按“作用、为什么这么写、容易追问点”三个层次回答。",
    )
    add_table(
        doc,
        ["文件/目录", "在项目中的作用", "为什么这样写", "面试追问重点"],
        [
            ["dataset_01.py", "定义 MF/Seq 训练和评估 Dataset，提供基础 DataLoader", "把矩阵分解与序列推荐的数据输入形式拆开，便于统一训练入口", "正负样本、序列 padding、候选集评估如何构造"],
            ["datasets/processed/*/meta.json", "保存数据集用户数、物品数、交互数和稀疏度", "实验结果必须能追溯数据规模和预处理条件", "MovieLens 与 Amazon 稀疏度差异为什么重要"],
            ["datasets/splits/*", "保存 leave-one-out 的 train/val/test 和 eval_candidates", "固定划分和候选集，保证不同方法可公平复现", "为什么不用每次评估重新采样负样本"],
            ["exp_01_fast_gpu/run_exp01_fast_gpu.py", "主实验脚本，包含 SASRec、FedMF、FedDIEG 和多种基线实现", "所有方法共享同一批量化训练与评估协议，减少实验设置差异", "FedDIEGModel、build_weighted_neighbors、update_memory_bank"],
            ["exp_02_fast_gpu/run_exp02_fast_gpu.py", "动态窗口实验脚本，按时间窗口评估兴趣漂移鲁棒性", "静态测试无法体现兴趣变化，需要 rolling windows 检验", "avg_window_ndcg、boundary_drop、adaptation_gain 含义"],
            ["exp_03_fast_gpu/run_exp03_fast_gpu.py", "消融实验脚本，比较 Ours-Plain、NoTemp、NoStruct、Full", "验证时间一致性和结构一致性是否真的带来稳定性收益", "avg_ecr 为什么能衡量图结构震荡"],
            ["exp_04_fast_gpu/run_exp04_fast_gpu.py", "新客户端冷启动实验脚本，实现原型语义锚定与适配过程", "模拟新用户只暴露少量交互时的推荐效果", "d_init、prototype memory mix、prototype score mix"],
            ["exp_05_sensitivity/run_exp05_sensitivity.py", "超参数敏感性分析脚本", "观察 K、rho、lambda_temp、lambda_struct 对效果和稳定性的影响", "参数调大为什么不一定更好"],
            ["figures/*", "保存论文图表脚本和敏感性/冷启动可视化结果", "让实验结果能生成可汇报、可复现的图表", "图中曲线和表格指标如何解释"],
        ],
        [1.55, 1.65, 1.7, 1.7],
    )

    add_heading(doc, "面试常见问题与回答模板", 1)
    qa_rows = [
        ["项目整体是做什么的？", "FedDIEG 是面向联邦推荐场景的动态序列推荐模型。它用 SASRec 建模本地用户行为序列，用服务器端兴趣记忆库和客户端动态兴趣图聚合相似用户表示，再通过一致性约束和原型锚定提升兴趣漂移和冷启动场景下的推荐效果。"],
        ["为什么要做联邦推荐？", "真实场景中用户行为数据往往涉及隐私，不能简单集中上传；同时不同平台、设备或边缘节点的数据天然分散。联邦推荐希望在不直接共享原始行为序列的前提下完成协同建模。我的实验是基于公开数据集的单机联邦模拟，每个用户模拟为一个客户端。"],
        ["为什么选择 SASRec 作为主干？", "SASRec 是经典序列推荐模型，使用 Transformer self-attention 捕捉历史物品之间的依赖关系。它比 FedMF 这类矩阵分解更能建模行为顺序，也比简单平均历史行为更适合表达短期兴趣变化。"],
        ["FedDIEG 和 FedSASRec 有什么区别？", "FedSASRec 主要是把 SASRec 放入联邦训练流程，本质上仍是本地序列建模加参数聚合。FedDIEG 额外维护客户端兴趣记忆库，构建动态兴趣图，并用图门控融合相似客户端信息，同时用时间/结构一致性约束提高动态场景稳定性。"],
        ["短期兴趣注意力解决什么问题？", "它解决只取最后位置或平均历史表示都不够稳的问题。注意力模块会给序列中不同位置分配权重，聚合出当前短期兴趣 summary，用于记忆库更新、邻居检索、图融合和最终打分。"],
        ["客户端动态兴趣图怎么构建？", "每个客户端有一个兴趣记忆向量。服务器对这些向量做归一化，用余弦相似度计算客户端间相似性，然后为每个客户端选择 Top-K 邻居，并用 temperature softmax 得到边权，低于 min_sim 的边会被过滤。"],
        ["图门控融合为什么必要？", "相似邻居不一定永远可靠，直接平均邻居表示可能引入负迁移。因此模型根据本地短期兴趣和图增强向量计算 gate，动态决定使用多少邻居信息；这样在本地信息更可靠时少引入外部信息，在邻居匹配时增强协同。"],
        ["时间一致性和结构一致性分别是什么？", "时间一致性约束同一客户端兴趣表示在相邻阶段不要剧烈跳变；结构一致性约束客户端动态图的边关系不要频繁震荡。前者稳定节点表示，后者稳定图结构，两者共同提升兴趣漂移场景下的鲁棒性。"],
        ["一致性约束会不会阻碍兴趣变化？", "会有这个风险，所以它们必须作为可调超参数。lambda_temp 和 lambda_struct 太大会让模型过于保守，太小又无法抑制噪声。敏感性分析就是为了观察推荐精度和稳定性之间的平衡。"],
        ["原型语义锚定是什么？", "它是冷启动机制。新客户端只有少量初始交互时，先得到一个粗糙兴趣表示，再在老客户端记忆库中检索相似原型，用这些稳定原型增强新客户端初始记忆和候选打分，帮助它更快进入合理兴趣区域。"],
        ["你如何模拟新客户端冷启动？", "实验中先划出一部分用户作为未见新客户端，模型只在老客户端上预训练。新客户端进入时只暴露 d_init=3 条交互，之后每轮逐步释放少量新交互，观察 0 到 10 轮适配过程中的 NDCG@10 和 HR@10。"],
        ["为什么要做 GPU 批量化？", "逐客户端模拟联邦训练速度很慢，GPU 利用率低。项目把同一轮采样到的客户端组成 microbatch，在 GPU 上批量完成序列编码、负采样、邻居检索和损失计算，既保留联邦客户端采样逻辑，又提高运行效率。"],
        ["负采样怎么做？", "对于每个有效序列位置，用真实下一个物品作为正样本，从用户未交互物品中随机采样负样本。训练目标提高正样本得分、降低负样本得分，代码中主要使用 BCE 形式的序列推荐损失。"],
        ["为什么评估用候选集？", "推荐系统全量物品排序成本高，很多序列推荐论文采用一个正样本加若干负样本的候选集评估。项目固定 eval_candidates，所有方法共享同一候选集，计算 HR@10、Recall@10 和 NDCG@10，保证对比公平。"],
        ["NDCG@10 和 HR@10 有什么区别？", "HR@10 只看真实物品是否进入前 10；NDCG@10 还考虑真实物品排在第几位，越靠前得分越高。因此 NDCG@10 更能反映推荐排序质量，我把它作为主要结果指标。"],
        ["项目的核心创新点是什么？", "第一是客户端动态兴趣图，在不共享原始行为的前提下建模相似客户端关系；第二是短期兴趣驱动的图门控融合，降低负迁移；第三是时间一致性、结构一致性和原型语义锚定，分别提升兴趣漂移稳定性和新客户端冷启动效果。"],
        ["你对比了哪些基线？", "包括 FedMF、FedSASRec、GPFedRec、FCUCR、TarDGR、F3CRec 和 DGFedRS。它们覆盖矩阵分解、序列推荐、个性化联邦推荐、检索增强、记忆增强和动态图推荐相关思路。"],
        ["主实验结果怎么讲？", "在统一协议下，FedDIEG 在四个数据集上相对最强基线取得提升，例如 Digital Music 的 NDCG@10 从 TarDGR 的 0.3247 提升到 0.3357，MovieLens-1M 从 0.4828 提升到 0.4934。需要稳妥说明不同数据集提升幅度不同。"],
        ["消融实验说明了什么？", "消融比较 Ours-Plain、NoTemp、NoStruct 和 Full。Arts/Crafts 多随机种子结果中，Full 将平均边变化率从约 0.9287 降到约 0.0129，稳定性得分从约 0.2228 提升到约 0.2662，说明一致性约束对控制动态图震荡有效。"],
        ["如果面试官质疑提升不大怎么办？", "可以回答推荐系统在强基线和公开数据集上几个百分点的相对提升已经有意义，更重要的是 FedDIEG 不只追求静态指标，还针对兴趣漂移和冷启动做了机制设计，消融和动态实验能说明模块确实改善稳定性和适应能力。"],
        ["真实部署和实验模拟有什么差别？", "实验是单机 GPU 上的联邦模拟，每个用户作为一个客户端。真实部署还需要通信协议、客户端调度、掉线处理、异构设备、安全聚合、差分隐私和攻击防御等工程模块。模型思想可以迁移，但工程链路需要补齐。"],
        ["项目有没有隐私风险？", "相比集中上传原始行为，风险更低；但上传模型参数或兴趣表示仍可能被反推。因此真实隐私敏感场景需要结合安全聚合、差分隐私、表示扰动或可信执行环境。本项目主要关注推荐性能和兴趣演化建模。"],
        ["这个项目最难的地方是什么？", "算法上难在把动态兴趣建模和联邦推荐约束结合起来，既要利用相似客户端信息，又不能访问原始数据；既要适应兴趣变化，又要避免动态图过度震荡。工程上难在统一多基线实验协议并实现高效 GPU 批量化训练。"],
        ["一句话怎么总结 FedDIEG？", "FedDIEG 是一个面向联邦推荐场景的动态兴趣演化建模方法，它用 SASRec 建模本地序列兴趣，用客户端动态兴趣图聚合相似用户表示，并通过一致性约束和原型锚定提升兴趣漂移与冷启动场景下的推荐效果。"],
    ]
    add_table(doc, ["面试问题", "回答模板"], qa_rows, [1.85, 4.75])

    add_heading(doc, "容易被追问的细节", 1)
    details = [
        "项目是联邦推荐模拟实验，不是真实多设备联邦部署。面试时不要夸大为已经完成端侧通信和安全聚合。",
        "每个用户模拟为一个客户端，原始交互序列不直接上传；服务器侧维护的是兴趣表示或模型相关信息。",
        "Top-K 邻居检索理论上有 O(N^2) 相似度计算压力，实验中使用 GPU 矩阵乘法和稀疏 Top-K；大规模部署可用 FAISS 或分桶近邻检索优化。",
        "动态图增强并非所有数据集都稳定领先，动态窗口实验中 Video Games 和 Arts/Crafts 与最强基线接近但略低，因此回答结果时要讲整体趋势和模块作用，不要说所有场景绝对领先。",
        "时间一致性和结构一致性是稳定性约束，不是越大越好。它们会带来精度与稳定性的权衡，需要通过 lambda_temp、lambda_struct 和 rho 调节。",
        "原型语义锚定的收益依赖老客户端记忆库质量。如果新客户端兴趣与老客户端分布差异大，原型迁移效果会受限。",
        "NDCG@10 是主要排序指标，HR@10/Recall@10 在 leave-one-out 候选集协议下含义接近，但 NDCG 会进一步考虑命中位置。",
        "FedDIEG 与传统协同过滤的联系是都利用相似用户兴趣；区别是 FedDIEG 在联邦设定下基于表示相似度做动态图聚合，而不是直接访问用户-物品交互矩阵。",
        "项目可复现性来自固定随机种子、固定数据划分、固定候选集、保存 config.json、round_metrics.csv 和 summary.csv，并对关键实验做多随机种子分析。",
        "后续优化可以从三方面展开：近似邻居检索降低成本，差分隐私/安全聚合增强隐私，使用多兴趣原型或多通道动态图提升表示能力。",
    ]
    for item in details:
        add_bullet(doc, item)

    doc.core_properties.title = "FedDIEG 面试资料"
    doc.core_properties.subject = "面向动态兴趣演化的联邦序列推荐模型"
    doc.core_properties.comments = "Generated from the uploaded interview-material template without modifying the reference file."
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT_DOCX))


if __name__ == "__main__":
    build_document()
